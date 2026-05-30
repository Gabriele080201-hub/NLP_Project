from __future__ import annotations

import logging
import os
from pathlib import Path
from typing import Dict, List, Optional
import openpyxl

from .schemas import Product

logger = logging.getLogger(__name__)

# --- Fallback Mock Data for Testing & Offline Execution ---
MOCK_CATALOG = [
    Product(
        code="FPN1T",
        description="FRUCHTPUREE MANGO 1kg GEFR. FRUITIERE",
        unit="BEG",
        package_size="1 kg",
        aliases=["mango puree", "mango piree", "fruchtpueree mango"],
    ),
    Product(
        code="OEG25",
        description="SONNENBLUMENOEL 10l BIG CHEF",
        unit="KAN",
        package_size="10 l",
        aliases=["sonnenblumenoel big chef", "sunflower oil big chef", "oel big chef"],
    ),
    Product(
        code="EI35F",
        description="EIER SPEZIAL 180er OX (L) 63-73g",
        unit="KRT",
        package_size="180 pcs",
        aliases=["180er eier l", "eggs l", "large eggs", "eier spezial l"],
    ),
    Product(
        code="KART02",
        description="KARTOFFELSALAT 2kg HAUSGEMACHT",
        unit="KRT",
        package_size="2 kg",
        aliases=["kartoffelsalat", "potato salad", "patate insalata"],
    ),
    Product(
        code="TOM10",
        description="TOMATENSAUCE PASSIERT 10kg",
        unit="EIM",
        package_size="10 kg",
        aliases=["tomatensauce", "passata", "tomato sauce"],
    ),
]

MOCK_CUSTOMER_HISTORY = {
    "CUST-DEMO": ["FPN1T", "OEG25", "EI35F"],
    "CUST-PIZZA": ["TOM10", "OEG25"],
    "CUST-CATERING": ["KART02", "EI35F", "FPN1T"],
}

# --- In-Memory Caches ---
CATALOG: List[Product] = list(MOCK_CATALOG)
CUSTOMER_HISTORY: Dict[str, List[str]] = dict(MOCK_CUSTOMER_HISTORY)


def load_all_data(data_dir: Optional[str] = None) -> bool:
    """
    Parses Schablone.xlsx and CompleteItemArchive.xlsx from the given data_dir
    and loads them into global memory catalog and customer history.
    
    If data_dir is not provided, reads from DATA_DIR environment variable.
    If the Excel files are not found or fail to load, falls back to the mock data.
    """
    global CATALOG, CUSTOMER_HISTORY
    
    if not data_dir:
        data_dir = os.getenv("DATA_DIR")
        
    if not data_dir:
        logger.warning("DATA_DIR env variable not set. Using fallback mock data.")
        return False
        
    data_path = Path(data_dir)
    archive_file = data_path / "2. Masterdata" / "CompleteItemArchive.xlsx"
    template_file = data_path / "2. Masterdata" / "Schablone.xlsx"
    
    if not archive_file.exists() or not template_file.exists():
        logger.warning(
            f"Excel files not found at: {archive_file} or {template_file}. "
            "Using fallback mock data."
        )
        return False
        
    try:
        logger.info(f"Loading master catalog from {archive_file}...")
        wb_archive = openpyxl.load_workbook(str(archive_file), read_only=True, data_only=True)
        sheet_archive = wb_archive["Tabelle1"]
        
        loaded_catalog: List[Product] = []
        seen_codes = set()
        
        # Format: ItemCode, DescriptionGerman, DescriptionItalian, UnitofMeasurement, PCPerUnit
        for row in sheet_archive.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            item_code = str(row[0]).strip()
            desc_de = str(row[1]).strip() if row[1] is not None else ""
            desc_it = str(row[2]).strip() if row[2] is not None else ""
            unit = str(row[3]).strip() if row[3] is not None else ""
            pc_per_unit = str(row[4]).strip() if row[4] is not None else "1"
            
            aliases = []
            if desc_it and desc_it != desc_de:
                aliases.append(desc_it)
                
            if item_code not in seen_codes:
                seen_codes.add(item_code)
                p = Product(
                    code=item_code,
                    description=desc_de,
                    unit=unit,
                    package_size=pc_per_unit,
                    aliases=aliases
                )
                loaded_catalog.append(p)
                
        logger.info(f"Successfully loaded {len(loaded_catalog)} items from catalog.")

        logger.info(f"Loading templates from {template_file}...")
        wb_template = openpyxl.load_workbook(str(template_file), read_only=True, data_only=True)
        sheet_template = wb_template["Tabelle1"]
        
        loaded_history: Dict[str, List[str]] = {}
        template_count = 0
        
        # Format: Customercode, ItemCode, DescriptionGerman, DescriptionItalian, Blocked, PCPerUnit, UnitofMeasurement
        for row in sheet_template.iter_rows(min_row=2, values_only=True):
            if not row or not row[0] or not row[1]:
                continue
            cust_code = str(row[0]).strip()
            item_code = str(row[1]).strip()
            desc_de = str(row[2]).strip() if row[2] is not None else ""
            desc_it = str(row[3]).strip() if row[3] is not None else ""
            blocked = row[4]
            pc_per_unit = str(row[5]).strip() if row[5] is not None else "1"
            unit = str(row[6]).strip() if row[6] is not None else ""
            
            # Skip blocked items
            if blocked == 1 or blocked == "1":
                continue
                
            if cust_code not in loaded_history:
                loaded_history[cust_code] = []
            loaded_history[cust_code].append(item_code)
            template_count += 1
            
            # Add to catalog if missing from the complete catalog
            if item_code not in seen_codes:
                seen_codes.add(item_code)
                aliases = []
                if desc_it and desc_it != desc_de:
                    aliases.append(desc_it)
                p = Product(
                    code=item_code,
                    description=desc_de,
                    unit=unit,
                    package_size=pc_per_unit,
                    aliases=aliases
                )
                loaded_catalog.append(p)

        logger.info(
            f"Successfully loaded templates for {len(loaded_history)} customers "
            f"({template_count} template records)."
        )

        # Update in-memory globals
        CATALOG.clear()
        CATALOG.extend(loaded_catalog)
        CUSTOMER_HISTORY.clear()
        CUSTOMER_HISTORY.update(loaded_history)
        
        return True

    except Exception as e:
        logger.error(f"Error loading Excel data files: {e}. Falling back to mock data.", exc_info=True)
        return False
